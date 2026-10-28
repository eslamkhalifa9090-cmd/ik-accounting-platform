"""Generates the migration that adds one IFRS-course level's questions to its level exam.

Run: python3 scripts/build_ifrs_questions.py <level number> <questions.xlsx> <migration number>
Example: python3 scripts/build_ifrs_questions.py 2 Ch02-Question-Bank-50.xlsx 0011
"""
import json, pathlib, sys
from openpyxl import load_workbook


def q(v):
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


SLUG = "ifrs"
level, xlsx, num = int(sys.argv[1]), sys.argv[2], sys.argv[3]
exam = f"(SELECT e.id FROM exams e JOIN courses c ON c.id = e.course_id WHERE c.slug = {q(SLUG)} AND e.kind = 'level' AND e.after_level = {level})"

out = [f"-- كورس المعايير الدولية IFRS: أسئلة امتحان المستوى {level} (اتولد من scripts/build_ifrs_questions.py)"]
# نفس أعمدة لصق الأدمن، والخانات الفاضية (أسئلة الصح والغلط) بتتشال، والسؤال مابيتكررش لو الملف اتنفذ تاني
ws = load_workbook(xlsx).worksheets[0]
n = 0
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row[0]:
        continue
    body, opts, ans, expl = row[0], [o for o in row[1:5] if o not in (None, "")], int(row[5]), row[6]
    assert 2 <= len(opts) <= 4 and 1 <= ans <= len(opts), body
    out.append("INSERT INTO questions (exam_id, body, options, answer, explanation) "
               f"SELECT {exam}, {q(body)}, {q(json.dumps(opts, ensure_ascii=False, separators=(",", ":")))}, {ans - 1}, {q(expl)} "
               f"WHERE NOT EXISTS (SELECT 1 FROM questions WHERE exam_id = {exam} AND body = {q(body)});")
    n += 1

path = pathlib.Path(__file__).resolve().parent.parent / "migrations" / f"{num}_ifrs_l{level}.sql"
path.write_text("\n".join(out) + "\n", encoding="utf-8")
print("wrote", path, n, "questions")
